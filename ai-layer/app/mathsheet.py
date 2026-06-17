"""Game math sheet exporter — the human-readable `.xlsx` deliverable a studio/operator/lab expects
(the readable face of the certification artifact; see artifact.py for the machine fingerprint).

Builds a multi-sheet workbook from a validated config + the engine's simulate results:
Summary · Symbols · Reels · Paytable · Win distribution. This is the first slice of the M4 GLI/BMM
exporter — it lays out exactly what a slot mathematician hands over when you "order math".
"""
from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

GOLD = "F5B942"
INK_HEAD = "1B2336"
INK_BAND = "10162A"
TITLE_FONT = Font(bold=True, size=15, color="FFFFFF")
SUB_FONT = Font(size=10, color="9AA4BF")
HEAD_FONT = Font(bold=True, size=10, color="FFFFFF")
KEY_FONT = Font(bold=True, color="C9D2EA")
GOLD_FONT = Font(bold=True, color="7A5A10")
HEAD_FILL = PatternFill("solid", fgColor=INK_HEAD)
GOLD_FILL = PatternFill("solid", fgColor=GOLD)
BAND_FILL = PatternFill("solid", fgColor=INK_BAND)
THIN = Side(style="thin", color="D7DEEC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _pct(x: Any) -> str:
    return f"{x * 100:.4f}%" if isinstance(x, (int, float)) else "—"


def _header_row(ws, row: int, cols: list[str]):
    for c, text in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=c, value=text)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="left", vertical="center")


def _widths(ws, widths: list[int]):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _summary_sheet(wb, config, results, config_hash, generated_at):
    ws = wb.active
    ws.title = "Summary"
    _widths(ws, [30, 46])
    ws.merge_cells("A1:B1")
    ws["A1"] = "SaS·Math — Game Math Sheet"
    ws["A1"].font = TITLE_FONT
    ws["A1"].fill = HEAD_FILL
    ws.merge_cells("A2:B2")
    ws["A2"] = config.get("game", {}).get("name", "Untitled")
    ws["A2"].font = SUB_FONT

    game = config.get("game", {})
    grid = game.get("grid", {})
    rows = [
        ("— GAME —", ""),
        ("Name", game.get("name", "")),
        ("Grid", f"{grid.get('reels', '?')} reels × {grid.get('rows', '?')} rows"),
        ("Bet units (RTP denominator)", game.get("bet_units", "")),
        ("Denomination scale", game.get("denomination_scale", "")),
        ("Win evaluation", config.get("line_eval", "")),
        ("Paylines", len(config.get("paylines", []))),
        ("Max-win cap", f"{config.get('max_win', {}).get('cap_x', '?')}×"),
        ("— RTP & VOLATILITY —", ""),
        ("Exact RTP (certified)", _pct(results.get("exact_rtp"))),
        ("  · line component", _pct(results.get("line_rtp"))),
        ("  · scatter component", _pct(results.get("scatter_rtp"))),
        ("Exact RTP (capped)", _pct(results.get("capped_rtp")) if results.get("capped_rtp") is not None else "numeric tail (MC)"),
        ("Monte-Carlo RTP", f"{_pct(results.get('mc_rtp'))}  ± {_pct(results.get('mc_ci95'))} (95%)"),
        ("Hit frequency", _pct(results.get("hit_frequency"))),
        ("Volatility (σ)", round(results.get("volatility", 0), 4)),
        ("Max win observed", f"{round(results.get('max_win_x', 0), 1)}×"),
        ("MC spins", results.get("spins", "")),
        ("— REPRODUCIBILITY —", ""),
        ("Exact line RTP (integer ratio)", f"{results.get('exact_line_ratio_num', '?')} / {results.get('exact_line_ratio_den', '?')}"),
        ("Enumerated combinations", results.get("exact_combos", "")),
        ("Config fingerprint", config_hash),
        ("Schema version", config.get("schema_version", "")),
        # Report the RNG the engine ACTUALLY uses, not the raw config string (they must never diverge).
        ("RNG", f"ChaCha8 — counter-based, namespaced (subsystem·spin·draw) · seed {config.get('rng', {}).get('master_seed', '')}"),
        ("Generated (UTC)", generated_at),
    ]
    r = 4
    for k, v in rows:
        kc = ws.cell(row=r, column=1, value=k)
        vc = ws.cell(row=r, column=2, value=v)
        if k.startswith("—"):
            kc.font = HEAD_FONT
            kc.fill = HEAD_FILL
            vc.fill = HEAD_FILL
        else:
            kc.font = KEY_FONT
            if k == "Exact RTP (certified)":
                vc.font = GOLD_FONT
                vc.fill = GOLD_FILL
        r += 1


def _symbols_sheet(wb, config):
    ws = wb.create_sheet("Symbols")
    _widths(ws, [8, 18, 12, 22])
    _header_row(ws, 1, ["ID", "Name", "Kind", "Substitutes for"])
    for i, sym in enumerate(config.get("symbols", []), start=2):
        ws.cell(row=i, column=1, value=sym.get("id"))
        ws.cell(row=i, column=2, value=sym.get("name"))
        ws.cell(row=i, column=3, value=sym.get("kind"))
        ws.cell(row=i, column=4, value=", ".join(map(str, sym.get("substitutes_for", []))) or "—")


def _reels_sheet(wb, config):
    ws = wb.create_sheet("Reels")
    reels = config.get("reels", {})
    names = {s["id"]: s["name"] for s in config.get("symbols", [])}
    if reels.get("strips"):
        _widths(ws, [10, 60])
        _header_row(ws, 1, ["Reel", "Strip (symbol ids, top→bottom cyclic)"])
        for i, strip in enumerate(reels["strips"], start=2):
            ws.cell(row=i, column=1, value=f"Reel {i - 1}")
            ws.cell(row=i, column=2, value=", ".join(str(x) for x in strip))
    elif reels.get("weights"):
        _widths(ws, [10, 8, 18, 12, 14])
        _header_row(ws, 1, ["Reel", "Sym id", "Symbol", "Weight", "Probability"])
        r = 2
        for ri, wmap in enumerate(reels["weights"], start=1):
            total = sum(int(v) for v in wmap.values()) or 1
            for sid, w in sorted(wmap.items(), key=lambda kv: int(kv[0])):
                ws.cell(row=r, column=1, value=f"Reel {ri}")
                ws.cell(row=r, column=2, value=int(sid))
                ws.cell(row=r, column=3, value=names.get(int(sid), f"#{sid}"))
                ws.cell(row=r, column=4, value=int(w))
                ws.cell(row=r, column=5, value=round(int(w) / total, 5))
                r += 1


def _paytable_sheet(wb, config):
    ws = wb.create_sheet("Paytable")
    names = {s["id"]: s["name"] for s in config.get("symbols", [])}
    counts = sorted({int(c) for e in config.get("paytable", []) for c in e.get("pays", {})})
    _widths(ws, [18, 8] + [10] * len(counts))
    _header_row(ws, 1, ["Symbol", "ID"] + [f"{c}-of-a-kind" for c in counts])
    for i, e in enumerate(config.get("paytable", []), start=2):
        ws.cell(row=i, column=1, value=names.get(e.get("symbol_id"), "?"))
        ws.cell(row=i, column=2, value=e.get("symbol_id"))
        for j, c in enumerate(counts, start=3):
            ws.cell(row=i, column=j, value=e.get("pays", {}).get(str(c), 0))


def _distribution_sheet(wb, results):
    ws = wb.create_sheet("Win distribution")
    _widths(ws, [18, 16, 16])
    _header_row(ws, 1, ["Win band (× bet)", "Probability", "MC count"])
    for i, b in enumerate(results.get("win_distribution", []), start=2):
        ws.cell(row=i, column=1, value=b.get("label"))
        ws.cell(row=i, column=2, value=round(b.get("probability", 0), 6))
        ws.cell(row=i, column=3, value=b.get("count"))


def build_math_sheet_xlsx(config: dict[str, Any], results: dict[str, Any], config_hash: str) -> bytes:
    generated_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    wb = Workbook()
    _summary_sheet(wb, config, results, config_hash, generated_at)
    _symbols_sheet(wb, config)
    _reels_sheet(wb, config)
    _paytable_sheet(wb, config)
    _distribution_sheet(wb, results)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
